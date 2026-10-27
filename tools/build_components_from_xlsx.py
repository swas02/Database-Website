"""
build_components_from_xlsx.py — builds every <bridge>.<sub|sup>.component.json from data/components.xlsx
(the sheets written by make_components_xlsx.py), then run reg.py to refresh the registry.

Column 1 is bridge_id; every other header is the dotted path of a value in the JSON file (a number in the
path = list position, 1-based). Empty cells are left out of the file. Girder segments are given as lengths
(girder_sections.longitudinal_segmentation_m.N.length_m); initial_m / end_m are added up from the support and the
segments are named support, intermediate, mid, intermediate, support.
"""

import json
from pathlib import Path

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent.parent
WORKBOOK = BASE_DIR / "data" / "components.xlsx"
DB_DIR = BASE_DIR / "project_database" / "PSC_vs_Steel_Short_Span_Bridges"
SHEETS = {"Steel_Sub": ("Sub_Str", "sub"), "Steel_Sup": ("Sup_Str", "sup"),
          "PSC_Sub": ("Sub_Str", "sub"), "PSC_Sup": ("Sup_Str", "sup")}
SEGMENT_NAMES = ("support", "intermediate", "mid", "intermediate", "support")


def insert(tree: dict, path: str, value):
    """Set tree["a"]["b"]["1"] = value for the path "a.b.1" (numbers become list positions later)."""
    keys = path.split(".")
    for k in keys[:-1]:
        tree = tree.setdefault(k, {})
    tree[keys[-1]] = value


def to_lists(node):
    """Dicts whose keys are 1..n become lists."""
    if not isinstance(node, dict):
        return node
    node = {k: to_lists(v) for k, v in node.items()}
    if node and all(k.isdigit() for k in node):
        return [node[k] for k in sorted(node, key=int)]
    return node


def build_segments(data: dict):
    sections = data.get("girder_sections") or {}
    lengths = sections.get("longitudinal_segmentation_m")
    if not lengths:
        return
    pos, segments = 0, []
    for name, seg in zip(SEGMENT_NAMES, lengths):
        length = seg.get("length_m") or 0
        segments.append({"section": name, "initial_m": round(pos, 6), "end_m": round(pos + length, 6)})
        pos += length
    sections["longitudinal_segmentation_m"] = segments


def main():
    wb = load_workbook(WORKBOOK, data_only=True)
    written = 0
    for sheet, (folder, kind) in SHEETS.items():
        rows = wb[sheet].iter_rows(values_only=True)
        header = next(rows)
        for row in rows:
            bridge_id = row[0]
            if not bridge_id:
                continue
            tree = {}
            for path, value in zip(header[1:], row[1:]):
                if value is not None and value != "":
                    insert(tree, path, value)
            data = to_lists(tree)
            build_segments(data)
            if bridge_id[-1] != sheet[0]:  # Steel_* sheets hold ..._S bridges, PSC_* sheets ..._P bridges
                print(f"skipped {sheet} {bridge_id}: wrong bridge type for this sheet")
                continue
            ids = [bridge_id]
            for bid in ids:
                bridge_dirs = list(DB_DIR.glob(f"*/*/{bid}"))
                target = bridge_dirs[0] / folder if bridge_dirs else None
                if not data or not target or not target.is_dir():
                    print(f"skipped {sheet} {bid}")
                    continue
                with open(target / f"{bid}.{kind}.component.json", "w", encoding="utf-8") as f:
                    json.dump(data, f, indent=2)
                    f.write(chr(10))
                written += 1
    print(f"wrote {written} component.json files")


if __name__ == "__main__":
    main()
