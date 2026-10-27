"""
make_components_xlsx.py — TEMPORARY helper. Builds data/components.xlsx, the sheet the
<bridge>.<sub|sup>.component.json files are meant to be created from.

Sheets: Steel_Sub, Steel_Sup, PSC_Sub, PSC_Sup. One row per bridge that has that part (from registry_flat.json).
Column 1 is bridge_id; every other header is the dotted path of a value in the JSON file, e.g.
  pile.count   pile_cap.dimensions_m.width_m   girder_sections.mid.web_depth_m
  girder_sections.longitudinal_segmentation_m.1.length_m      (a number in the path = list position, 1-based)
Girder segments are asked as a length (length_m), not as initial/end: the file builder works out
initial_m / end_m by adding the lengths up from the support, and names the segments in the fixed
order support, intermediate, mid, intermediate, support (true for all 48 steel bridges), so the sheet has no section column.
A girder section whose length_per_girder_m is 0 does not exist, so its flange / web size cells are greyed out
(and the file builder must write them as 0).
Columns of the same girder segment sit side by side and their headers (support / intermediate / mid, incl. their length columns) share a colour.
Steel rows are filled from archive/steel_data_v2.json. PSC rows are empty, to be filled in.
"""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from split_steel_data import STEEL_DATA, split

BASE_DIR = Path(__file__).resolve().parent.parent
REGISTRY = BASE_DIR / "registry_flat.json"
OUTPUT = BASE_DIR / "data" / "components.xlsx"

# PSC bridges have no steel girders; these are the superstructure fields the viewer can use for them
PSC_SUP_COLUMNS = ["total_width_m", "deck_slab.depth_m", "girder_spacing_m",
                   "material_grades.rebar", "material_grades.concrete_substructure_and_deck_slab"]


def flatten(value, prefix="") -> dict:
    """Nested dict/list -> {"a.b.1.c": leaf}. List positions are 1-based."""
    if isinstance(value, dict):
        items = value.items()
    elif isinstance(value, list):
        items = ((i + 1, v) for i, v in enumerate(value))
    else:
        return {prefix: value}
    out = {}
    for k, v in items:
        out.update(flatten(v, f"{prefix}.{k}" if prefix else str(k)))
    return out


def segments_as_lengths(data: dict) -> dict:
    """Replace each girder segment's section / initial_m / end_m by its length_m."""
    segs = (data.get("girder_sections") or {}).get("longitudinal_segmentation_m")
    if segs:
        data["girder_sections"]["longitudinal_segmentation_m"] = [
            {"length_m": round(g["end_m"] - g["initial_m"], 6)} for g in segs]
    return data


SECTIONS = ("mid", "intermediate", "support")
SECTION_DIMS = ("top_flange_width_m", "top_flange_thickness_m", "web_depth_m", "web_thickness_m",
                "bottom_flange_width_m", "bottom_flange_thickness_m")


def grey_unused_sections(ws, columns):
    """Grey out a section's flange / web cells while its length_per_girder_m is 0."""
    grey = PatternFill("solid", bgColor="D9D9D9", fgColor="D9D9D9")
    for sec in SECTIONS:
        prefix = f"girder_sections.{sec}."
        if prefix + "length_per_girder_m" not in columns:
            continue
        len_col = get_column_letter(columns.index(prefix + "length_per_girder_m") + 2)
        for dim in SECTION_DIMS:
            col = get_column_letter(columns.index(prefix + dim) + 2)
            ws.conditional_formatting.add(
                f"{col}2:{col}{ws.max_row}",
                FormulaRule(formula=[f'AND(${len_col}2<>"",${len_col}2=0)'], fill=grey, font=Font(color="808080")))


# Header colour per girder segment: its size columns and the length columns of its positions share it
SEGMENT_COLOURS = {"support": "F8CBAD", "intermediate": "FFE699", "mid": "C6E0B4", "end_diaphragm": "BDD7EE"}
SEGMENT_POSITIONS = {"1": "support", "2": "intermediate", "3": "mid", "4": "intermediate", "5": "support"}


def segment_of(column: str) -> str | None:
    parts = column.split(".")
    if parts[0] != "girder_sections":
        return None
    if parts[1] == "longitudinal_segmentation_m":
        return SEGMENT_POSITIONS.get(parts[2])
    return parts[1] if parts[1] in SEGMENT_COLOURS else None


def group_by_segment(columns: list) -> list:
    """Move each girder segment's columns together: size columns, length_per_girder_m, then its length_m positions."""
    order = list(SEGMENT_COLOURS)  # support, intermediate, mid, end_diaphragm
    idx = [i for i, c in enumerate(columns) if segment_of(c)]
    if not idx:
        return columns
    # stable sort keeps the size columns in file order; the length_m positions come last within their segment
    block = sorted((columns[i] for i in idx),
                   key=lambda c: (order.index(segment_of(c)), "longitudinal_segmentation_m" in c))
    out = list(columns)
    for i, c in zip(idx, block):
        out[i] = c
    return out


def write_sheet(wb, title, columns, rows):
    ws = wb.create_sheet(title)
    columns = group_by_segment(columns)
    ws.append(["bridge_id", *columns])
    for bridge_id, values in rows:
        ws.append([bridge_id, *(values.get(c) for c in columns)])
    head = PatternFill("solid", fgColor="DDEBF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = head
        cell.alignment = Alignment(vertical="center")
    for i, name in enumerate(columns, start=2):
        if segment_of(name):  # whole column, header and rows
            fill = PatternFill("solid", fgColor=SEGMENT_COLOURS[segment_of(name)])
            for row in range(1, ws.max_row + 1):
                ws.cell(row, i).fill = fill
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 16
    for i, name in enumerate(columns, start=2):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(name) + 4)
    ws.auto_filter.ref = ws.dimensions
    grey_unused_sections(ws, columns)


def main():
    registry = json.load(open(REGISTRY, encoding="utf-8"))
    steel = {e["bridge_id"]: e for e in json.load(open(STEEL_DATA, encoding="utf-8-sig"))}
    flat = {bid: {part: flatten(segments_as_lengths(data)) for part, data in split(spec).items()} for bid, spec in steel.items()}

    wb = Workbook()
    wb.remove(wb.active)
    for bridge_type in ("Steel", "PSC"):
        for part, key, dirname in (("Sub", "sub", "Sub_Str"), ("Sup", "sup", "Sup_Str")):
            ids = [b["bridge_id"] for b in registry if b["bridge_type"] == bridge_type and b[key]]
            if bridge_type == "Steel":
                columns = list(dict.fromkeys(c for bid in ids for c in flat.get(bid, {}).get(dirname, {})))
            elif key == "sub":  # same substructure fields as steel
                columns = list(dict.fromkeys(c for d in flat.values() for c in d["Sub_Str"]))
            else:
                columns = PSC_SUP_COLUMNS
            rows = [(bid, flat.get(bid, {}).get(dirname, {}) if bridge_type == "Steel" else {}) for bid in ids]
            write_sheet(wb, f"{bridge_type}_{part}", columns, rows)
            print(f"{bridge_type}_{part}: {len(rows)} bridges, {len(columns)} columns")
    wb.save(OUTPUT)
    print(f"wrote {OUTPUT}")


if __name__ == "__main__":
    main()
