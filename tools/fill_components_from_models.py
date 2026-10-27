"""
fill_components_from_models.py — makes data/components_from_model.xlsx: a copy of data/components.xlsx with every value
cleared (headers and the bridge_id column are kept), then filled with sizes measured on each bridge's model.obj.

Only columns that can be measured from the geometry are filled; everything else (grades, reinforcement %, PCC depth,
pedestal size, girder section sizes) stays empty. Axes of the models: x along the bridge, y up, z across the bridge.
Bridge parts are read from <bridge>/Sub_Str/model.obj and <bridge>/Sup_Str/model.obj, mesh groups by their `g` name.
The original components.xlsx is never changed.
"""

import shutil
import statistics
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent.parent
SOURCE = BASE_DIR / "data" / "components.xlsx"
OUTPUT = BASE_DIR / "data" / "components_from_model.xlsx"
DB_DIR = BASE_DIR / "project_database" / "PSC_vs_Steel_Short_Span_Bridges"
SHEET_FOLDER = {"Steel_Sub": "Sub_Str", "Steel_Sup": "Sup_Str", "PSC_Sub": "Sub_Str", "PSC_Sup": "Sup_Str"}
WELD_DIGITS = 3  # vertices closer than 1 mm are one point


def norm(name: str) -> str:
    return name.lower().replace(" ", "_")


def load_obj(path: Path) -> dict:
    """{normalised group name: [connected pieces]}; a piece is a list of 3 (min, max) pairs (x, y, z) plus its vertices."""
    verts, faces, group = [], defaultdict(list), None
    for line in open(path, encoding="utf-8", errors="ignore"):
        if line.startswith("v "):
            verts.append(tuple(float(c) for c in line.split()[1:4]))
        elif line.startswith("g "):
            group = norm(line[2:].strip())
        elif line.startswith("f "):
            faces[group].append([int(t.split("/")[0]) - 1 for t in line.split()[1:]])
    welded, key = [], {}
    for v in verts:
        welded.append(key.setdefault(tuple(round(c, WELD_DIGITS) for c in v), len(key)))
    points = list(key)
    out = {}
    for g, fs in faces.items():
        parent = {}

        def find(a):
            while parent.setdefault(a, a) != a:
                parent[a] = parent[parent[a]]
                a = parent[a]
            return a

        for f in fs:
            ids = [welded[i] for i in f]
            for i in ids[1:]:
                parent[find(i)] = find(ids[0])
        pieces = defaultdict(list)
        for i in list(parent):
            pieces[find(i)].append(points[i])
        out[g] = sorted((Piece(p) for p in pieces.values()), key=lambda p: (p.c(0), p.c(2)))
    return out


class Piece:
    def __init__(self, pts):
        self.pts = pts
        self.lo = [min(p[a] for p in pts) for a in range(3)]
        self.hi = [max(p[a] for p in pts) for a in range(3)]

    def size(self, a):
        return self.hi[a] - self.lo[a]

    def c(self, a):
        return (self.lo[a] + self.hi[a]) / 2


def find_group(groups: dict, *names, prefix=False) -> list:
    """Pieces of every group whose name matches one of names."""
    out = []
    for g, pieces in groups.items():
        if any(g == n or (prefix and g.startswith(n)) for n in names):
            out += pieces
    return out


def span(pieces, a):
    return max(p.hi[a] for p in pieces) - min(p.lo[a] for p in pieces) if pieces else None


def spacing(values, tol=0.05):
    """Median gap between the distinct values (values closer than tol are the same position)."""
    pos = []
    for v in sorted(values):
        if not pos or v - pos[-1] > tol:
            pos.append(v)
    gaps = [b - a for a, b in zip(pos, pos[1:])]
    return statistics.median(gaps) if gaps else None, len(pos)


def girder_centres(girders):
    """z centre of every girder (a girder is made of several pieces side by side at the same z)."""
    groups = []
    for p in sorted(girders, key=lambda p: p.c(2)):
        if groups and p.c(2) - groups[-1][-1].c(2) < 1.0:
            groups[-1].append(p)
        else:
            groups.append([p])
    return [(min(q.lo[2] for q in g) + max(q.hi[2] for q in g)) / 2 for g in groups]


def bracing_section(piece):
    """Angle leg sizes (mm) and thickness (mm) of a diagonal brace: its 6-point L cross-section."""
    dy, dz = piece.size(1), piece.size(2)
    n = (dz ** 2 + dy ** 2) ** 0.5
    if n == 0:
        return None
    ny, nz = -dz / n, dy / n  # in-plane direction perpendicular to the brace axis
    pts = {(round(p[0], 3), round(p[1] * ny + p[2] * nz, 3)) for p in piece.pts}
    # 12 welded points = 6 per end; the 6 section points are those at one end of the axis
    ay, az = dy / n, dz / n
    ends = defaultdict(set)
    for p in piece.pts:
        ends[round(p[1] * ay + p[2] * az, 2)].add((round(p[0], 3), round(p[1] * ny + p[2] * nz, 3)))
    ends = [e for e in ends.values() if len(e) == 6]
    if not ends:
        return None
    sec = sorted(ends[0])
    us, vs = sorted({s[0] for s in sec}), sorted({s[1] for s in sec})
    leg1, leg2 = us[-1] - us[0], vs[-1] - vs[0]
    t = min(b - a for a, b in zip(us, us[1:])) if len(us) > 1 else None
    if not t:
        return None
    return round(leg1 * 1000), round(leg2 * 1000), round(t * 1000)


def sup_values(groups: dict, is_steel: bool) -> dict:
    out = {}
    deck = find_group(groups, "deck_slab")
    if deck:
        out["total_width_m"] = span(deck, 2)
        out["deck_slab.depth_m"] = span(deck, 1)
    girders = find_group(groups, "girder", prefix=True)
    centres = girder_centres(girders) if girders else []
    if len(centres) > 1:
        out["girder_spacing_m"] = statistics.median(b - a for a, b in zip(centres, centres[1:]))
    if not is_steel:
        return out
    chords = find_group(groups, "chord") or find_group(groups, "diagonal_brace")
    if chords:
        gap, count = spacing([p.c(0) for p in chords])
        out["cross_bracing.count"], out["cross_bracing.spacing_m"] = count, gap
    braces = find_group(groups, "diagonal_brace")
    section = bracing_section(braces[0]) if braces else None
    if section:
        leg1, leg2, t = section
        out["cross_bracing.section.raw"] = f"{leg1} x {leg2} x {t}"
        out["cross_bracing.section.leg1_mm"], out["cross_bracing.section.leg2_mm"] = leg1, leg2
        out["cross_bracing.section.thickness_mm"] = t
    for group, col in (("intermediate_stiffener", "transverse_stiffeners_per_girder_side"),
                       ("end_stiffener", "bearing_stiffeners_per_girder_side")):
        st = find_group(groups, group)
        if st:
            out[f"{col}.width_m"] = statistics.median(p.size(2) for p in st)
            out[f"{col}.thickness_m"] = statistics.median(p.size(0) for p in st)
            out[f"{col}.spacing_m"] = spacing([p.c(0) for p in st], tol=0.02)[0]
    return out


def number(v):
    return int(v) if float(v).is_integer() else v


def sub_values(groups: dict) -> dict:
    out = {}
    piles, caps = find_group(groups, "pile"), find_group(groups, "pile_cap")
    if piles:
        out["pile.count"] = len(piles) // max(len(caps), 1)
        out["pile.depth_m"] = statistics.median(p.size(1) for p in piles)
        out["pile.diameter_m"] = statistics.median((p.size(0) + p.size(2)) / 2 for p in piles)
    if caps:
        cap = caps[0]
        out["pile_cap.depth_m"] = cap.size(1)
        out["pile_cap.dimensions_m.length_m"], out["pile_cap.dimensions_m.width_m"] = cap.size(0), cap.size(2)
        out["pile_cap.dimensions_m.raw"] = f"{number(round(cap.size(2), 3))} X {number(round(cap.size(0), 3))}"
    piers = find_group(groups, "pier", "pier_shaft")
    if piers:
        out["pier.height_m"] = statistics.median(p.size(1) for p in piers)
        out["pier.diameter_m"] = statistics.median((p.size(0) + p.size(2)) / 2 for p in piers)
    pier_caps = find_group(groups, "pier_cap")
    if pier_caps:
        cap = pier_caps[0]
        deepest = [p for p in cap.pts if abs(p[1] - cap.lo[1]) < 0.005]
        end = [p[1] for p in cap.pts if abs(p[2] - cap.lo[2]) < 0.005]
        mid = max(p[2] for p in deepest) - min(p[2] for p in deepest)
        out["pier_cap.width_m"] = cap.size(0)
        out["pier_cap.mid_depth_m"] = cap.size(1)
        out["pier_cap.end_depth_m"] = max(end) - min(end)
        out["pier_cap.mid_section_length_m"] = mid
        out["pier_cap.taper_section_length_m"] = (cap.size(2) - mid) / 2
    return out


def main():
    shutil.copy(SOURCE, OUTPUT)
    wb = load_workbook(OUTPUT)
    filled = missing = 0
    for ws in wb:
        folder, is_sup, is_steel = SHEET_FOLDER[ws.title], ws.title.endswith("Sup"), ws.title.startswith("Steel")
        header = [c.value for c in ws[1]]
        for row in ws.iter_rows(min_row=2, min_col=2):  # clear all values, keep bridge_id and headers
            for cell in row:
                cell.value = None
        for r in range(2, ws.max_row + 1):
            bridge_id = ws.cell(r, 1).value
            dirs = list(DB_DIR.glob(f"*/*/{bridge_id}/{folder}/model.obj")) if bridge_id else []
            if not dirs:
                missing += 1
                print(f"no model: {ws.title} {bridge_id}")
                continue
            groups = load_obj(dirs[0])
            values = sup_values(groups, is_steel) if is_sup else sub_values(groups)
            for col, name in enumerate(header[1:], start=2):
                v = values.get(name)
                if v is not None:
                    ws.cell(r, col).value = round(v, 3) if isinstance(v, float) else v
                    filled += 1
    wb.save(OUTPUT)
    print(f"wrote {OUTPUT}: {filled} values filled, {missing} rows without a model")


if __name__ == "__main__":
    main()
