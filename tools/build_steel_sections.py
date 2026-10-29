"""
build_steel_sections.py — applies the steel-superstr sheet to every steel <bridge>.sup.component.json.

The sheet is the source of truth for the girder sections (support = "End Girder", intermediate = "Inter Girder",
mid = "Mid Girder"), the longitudinal segmentation, cross bracing, girder spacing, slab thickness and material grades.
Keys the sheet does not cover (stiffeners, total width, ...) are left as they are. npm run components runs it after
build_components_from_xlsx.py, which rewrites the component files from data/components.xlsx.

    python tools/build_steel_sections.py data/steel-superstr.xlsx

The sheet is never modified.
"""

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent.parent
STEEL_DIR = BASE_DIR / "project_database" / "PSC_vs_Steel_Short_Span_Bridges" / "Steel"
# component.json zone -> sheet column prefix
ZONES = {"support": "End", "intermediate": "Inter", "mid": "Mid"}
# component.json key -> sheet column suffix
FIELDS = {
    "top_flange_width_m": "B1(m)",
    "top_flange_thickness_m": "tf1(m)",
    "web_depth_m": "Hw(m)",
    "web_thickness_m": "tw(m)",
    "bottom_flange_width_m": "B2(m)",
    "bottom_flange_thickness_m": "tf2(m)",
}
SEGMENTS = [("support", "End 1(m)"), ("intermediate", "Intermediate 1(m)"), ("mid", "Mid(m)"),
            ("intermediate", "Intermediate 2(m)"), ("support", "End 2(m)")]


def num(v, default=0):
    """Sheet number, whole values kept as ints so the JSON reads 16, not 16.0."""
    if v in (None, ""):
        return default
    v = round(float(v), 4)
    return int(v) if v == int(v) else v


def section(row, zone):
    prefix = ZONES[zone]
    return {k: num(row.get(f"{prefix} Girder {col}")) for k, col in FIELDS.items()}


def bracing(raw):
    """'ISA 90x90x10' -> section dict as the component files hold it."""
    legs = [int(n) for n in re.findall(r"\d+", raw or "")]
    if len(legs) != 3:
        return None
    return {"raw": f"{legs[0]} x {legs[1]} x {legs[2]}", "leg1_mm": legs[0], "leg2_mm": legs[1], "thickness_mm": legs[2]}


def main(xlsx):
    ws = load_workbook(xlsx, data_only=True).active
    header, *rows = ws.iter_rows(values_only=True)
    written, problems = 0, []
    for values in rows:
        row = dict(zip(header, values))
        bridge_id = str(row["FileName"]).removesuffix("_Super")
        target = STEEL_DIR / row["Location"] / bridge_id / "Sup_Str" / f"{bridge_id}.sup.component.json"
        if not target.exists():
            problems.append(f"{bridge_id}: no {target.name}")
            continue
        data = json.loads(target.read_text(encoding="utf-8"))
        gs = data.setdefault("girder_sections", {})

        # the end diaphragm repeats the mid girder section; keep it that way unless it was set apart
        old_mid = {k: gs.get("mid", {}).get(k) for k in FIELDS}
        old_dia = gs.get("end_diaphragm")
        diaphragm_follows_mid = old_dia is not None and all(old_dia.get(k) == old_mid[k] for k in FIELDS)

        lengths = {}
        for zone, col in SEGMENTS:
            lengths[zone] = lengths.get(zone, 0) + num(row.get(col))
        for zone in ZONES:
            gs[zone] = {**section(row, zone), "length_per_girder_m": num(lengths.get(zone, 0))}
        if diaphragm_follows_mid:
            gs["end_diaphragm"] = section(row, "mid")

        segs, at = [], 0.0
        for zone, col in SEGMENTS:
            length = num(row.get(col))
            segs.append({"section": zone, "initial_m": num(at), "end_m": num(at + length)})
            at += length
        gs["longitudinal_segmentation_m"] = segs
        if abs(at - num(row.get("Total Length"))) > 0.01 or row.get("Verify length") != "Ok":
            problems.append(f"{bridge_id}: segments add up to {at:.3f} m, sheet total {row.get('Total Length')}")

        section_raw = bracing(row.get("Bracing"))
        if section_raw:
            data["cross_bracing"] = {
                "spacing_m": num(row.get("Bracing Spacing 1(m)")),
                "count": int(num(row.get("Number of Bracing Sets"))),
                "section": section_raw,
            }
        else:
            problems.append(f"{bridge_id}: bracing '{row.get('Bracing')}' not parsed")
        data["girder_spacing_m"] = num(row.get("Unique Girder Spacing 1(m)"))
        data.setdefault("deck_slab", {})["depth_m"] = num(row.get("Slab Thickness(m)"))
        grades = data.setdefault("material_grades", {})
        grades["structural_steel"] = row["Material 1"]
        grades["cross_bracing"] = row["Material 2"]
        grades["concrete_deck_slab"] = row["Material 3"]

        target.write_text(json.dumps(data, indent=2) + "\n", encoding="utf-8")
        written += 1
    print(f"updated {written} component files")
    for p in problems:
        print("  !", p)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
