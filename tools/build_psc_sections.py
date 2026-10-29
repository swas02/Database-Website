"""
build_psc_sections.py — adds the PSC girder cross-section data to every PSC <bridge>.sup.component.json.

Section depths and flange half-widths come from the psc-super-data sheet (HR1..HR5 = top flange, top haunch, web,
bottom haunch, bottom flange depth; BL2+BR2 = top flange width; BL4+BR4 = bottom flange width). The sheet has no web
thickness, so it is read from the model.obj: the narrowest girder width between the haunches. npm run components runs
it after build_components_from_xlsx.py, which rewrites the component files and drops these keys.

    python tools/build_psc_sections.py data/psc-super-data.xlsx

Adds "psc_girder" (sections per zone, girder and tendon counts) and "psc_diaphragm" (end diaphragm rectangle).
The sheet is never modified; diaphragm sizes typed in mm (1475) are converted to m.
"""

import json
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent.parent
PSC_DIR = BASE_DIR / "project_database" / "PSC_vs_Steel_Short_Span_Bridges" / "PSC"
# zone -> (sheet column prefix, OBJ group written by the converter)
ZONES = {
    "support": ("Support Section", "Girder_(support_section)"),
    "transition": ("{span}m_Support>Mid", "Girder_(Transition_Zone)"),
    "mid": ("Mid-Section", "Girder_(mid-portion)"),
}
LAYERS = ("HR1", "HR2", "HR3", "HR4", "HR5", "BL2", "BL4", "BR2", "BR4")


def num(v):
    return float(v) if v not in (None, "") else None


def metres(v):
    """Sizes are in m; a few cells hold mm (1475, 1700)."""
    v = num(v)
    return None if v is None else (v / 1000 if v > 10 else v)


def group_points(obj_path: Path):
    """Vertices of each OBJ group as {group: [(x, y, z), ...]}."""
    verts, faces, groups, cur = [], defaultdict(set), {}, None
    for line in obj_path.read_text().splitlines():
        if line.startswith("v "):
            verts.append(tuple(map(float, line.split()[1:4])))
        elif line.startswith("g "):
            cur = line[2:].strip()
        elif line.startswith("f ") and cur:
            faces[cur].update(int(t.split("/")[0]) - 1 for t in line.split()[1:])
    for g, ids in faces.items():
        groups[g] = [verts[i] for i in ids]
    return groups


def web_thickness(points, hr1, hr2, hr3, centre=False):
    """Narrowest width of the first girder between the top and bottom haunch (None if no vertices there).
    Measured at the group's first slice, or at the middle of the group for the mid zone, whose first slice can be a flare."""
    top = max(p[1] for p in points)
    xs = sorted({p[0] for p in points})
    x0 = xs[len(xs) // 2] if centre else xs[0]
    hi, lo = top - hr1 - hr2 + 2e-3, top - hr1 - hr2 - hr3 - 2e-3
    levels = defaultdict(set)
    for x, y, z in points:
        if abs(x - x0) < 1e-3 and lo <= y <= hi:
            levels[round(y, 3)].add(round(z, 3))
    widths = [zs[1] - zs[0] for zs in map(sorted, levels.values()) if len(zs) >= 2 and zs[1] - zs[0] > 0.05]
    return round(min(widths), 3) if widths else None


def build_zone(row, prefix, points, problems, label, centre=False):
    v = {k: metres(row.get(f"{prefix}_{k}")) for k in LAYERS}
    if any(v[k] is None for k in LAYERS):
        return None
    depth = sum(v[k] for k in ("HR1", "HR2", "HR3", "HR4", "HR5"))
    web = web_thickness(points, v["HR1"], v["HR2"], v["HR3"], centre) if points else None
    if points:
        got = max(p[1] for p in points) - min(p[1] for p in points)
        if abs(got - depth) > 0.005:
            problems.append(f"{label}: sheet depth {depth:.3f} m, model {got:.3f} m")
    return {
        "top_flange_width_m": round(v["BL2"] + v["BR2"], 4),
        "top_flange_thickness_m": v["HR1"],
        "top_haunch_depth_m": v["HR2"],
        "web_depth_m": v["HR3"],
        "web_thickness_m": web,
        "bottom_haunch_depth_m": v["HR4"],
        "bottom_flange_width_m": round(v["BL4"] + v["BR4"], 4),
        "bottom_flange_thickness_m": v["HR5"],
        "total_depth_m": round(depth, 4),
    }


def main(xlsx):
    ws = load_workbook(xlsx, data_only=True).active
    header, *rows = ws.iter_rows(values_only=True)
    written, problems = 0, []
    for values in rows:
        row = dict(zip(header, values))
        bridge_id = str(row["FileName"]).removesuffix("_Super")
        folder = PSC_DIR / row["Location"] / bridge_id / "Sup_Str"
        target = folder / f"{bridge_id}.sup.component.json"
        if not target.exists():
            problems.append(f"{bridge_id}: no {target.name}")
            continue
        obj = folder / "model.obj"
        groups = group_points(obj) if obj.exists() else {}

        sections = {}
        for zone, (prefix, group) in ZONES.items():
            sec = build_zone(row, prefix.format(span=int(row["Span(m)"])), groups.get(group), problems, f"{bridge_id} {zone}", centre=zone == "mid")
            if sec:
                sections[zone] = sec
        if not sections:
            problems.append(f"{bridge_id}: no section data")
            continue
        # the transition zone flares from the support end block to the mid web, which is the model's thin end
        mid_web = (sections.get("mid") or {}).get("web_thickness_m")
        end_web = (sections.get("support") or {}).get("web_thickness_m")
        if "transition" in sections and mid_web:
            sections["transition"]["web_thickness_m"] = mid_web
            if end_web and end_web != mid_web:
                sections["transition"]["web_thickness_support_end_m"] = end_web
        for sec in sections.values():
            sec["web_thickness_m"] = sec["web_thickness_m"] or mid_web

        data = json.loads(target.read_text())
        data["psc_girder"] = {
            "girder_count": int(row["Num_Girders"]),
            "tendon": row["Tendon_Name"],
            "tendon_count": int(row["Num_Tendons"]),
            "tendons_per_girder": int(row["Tendons_Per_Girder"]),
            "sections": sections,
        }
        h, b = metres(row["Diaphragm_H"]), metres(row["Diaphragm_B"])
        if h and b:
            data["psc_diaphragm"] = {"height_m": h, "width_m": b}
        target.write_text(json.dumps(data, indent=2) + "\n")
        written += 1
    print(f"updated {written} component files")
    for p in problems:
        print("  !", p)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
